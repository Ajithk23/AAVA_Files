from __future__ import annotations

from copy import copy
from datetime import date, datetime, time, timedelta
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment


BASE_DATE = date.today()
TEMPLATE_PATH = Path(__file__).with_name("Sample template.xlsx")
OUTPUT_PATH = TEMPLATE_PATH.with_name(
    f"NICE_Smoke_Testing_{BASE_DATE.isoformat()}.xlsx"
)


def fmt_date(value: date) -> str:
    return value.strftime("%Y-%m-%d")


def fmt_dt(value: datetime) -> str:
    return value.strftime("%Y-%m-%d %H:%M ET")


CALL_START = datetime.combine(BASE_DATE, time(11, 0))
CALL_END = CALL_START + timedelta(minutes=6)


TEST_CASES = [
    {
        "scenario": "NICE Smoke - Escalate to Live Agent Transfer",
        "description": "Verify that when escalation to a live agent is initiated from a supported voice flow, the interaction is transferred successfully into NICE CXone and reaches the intended live agent path.",
        "preconditions": "NICE CXone escalation path is configured; test agent logged in and available.",
        "steps": "1. Start an interaction that supports escalation to a live agent\n2. Choose the escalate-to-live-agent option\n3. Observe whether the interaction is handed off to NICE CXone\n4. Confirm the call lands in the expected agent or queue",
        "expected": "The escalation transfers successfully to NICE CXone, the call is routed to the correct live-agent queue or agent, and the interaction context is preserved.",
        "comments": f"Test Data: Today={fmt_date(BASE_DATE)}; escalation_source=Virtual Agent; target_queue=Live Agent Support; agent_state=Available; transfer_time={fmt_dt(CALL_START)}.",
    },
    {
        "scenario": "NICE Smoke - Queue Counter Validation",
        "description": "Verify NICE CXone displays correct queue counters for digital, inbound voice, voicemail, and work item channels.",
        "preconditions": "Agent has access to NICE CXone queue dashboard or agent panel; counters configured and visible.",
        "steps": "1. Open NICE CXone queue or agent panel\n2. Review counters for digital items\n3. Review counters for inbound voice\n4. Review counters for voicemail\n5. Review counters for work items\n6. Compare counts with current queue state",
        "expected": "Queue counters for digital, inbound voice, voicemail, and work items are visible and reflect the expected values for each channel.",
        "comments": f"Test Data: Today={fmt_date(BASE_DATE)}; queues_checked=Digital,Inbound Voice,Voicemail,Work Item; agent_id=agent_cx_smoke_01; validation_time={fmt_dt(CALL_START)}.",
    },
    {
        "scenario": "NICE Smoke - Mute Control During Call",
        "description": "Verify the mute function works correctly during an active NICE CXone call and can be toggled without dropping the call.",
        "preconditions": "Agent is logged in to NICE CXone and connected to an active call.",
        "steps": "1. Accept or place a test call\n2. Select the mute control\n3. Confirm microphone audio is muted\n4. Select unmute\n5. Confirm audio resumes normally",
        "expected": "Mute and unmute work correctly during the call, and the call remains connected throughout the action.",
        "comments": f"Test Data: active_call_id=SMOKE-CALL-03; call_start={fmt_dt(CALL_START)}; expected_call_state=Connected.",
    },
    {
        "scenario": "NICE Smoke - Transfer Call Option During Call",
        "description": "Verify the transfer option is available and functional while the agent is on an active call in NICE CXone.",
        "preconditions": "Agent is on an active call and has permission to transfer calls.",
        "steps": "1. Accept or place a test call\n2. Open the transfer option during the call\n3. Select a target agent or queue\n4. Complete the transfer\n5. Verify the call reaches the selected destination",
        "expected": "The transfer option works during the call, and the call is routed successfully to the chosen destination without unexpected disconnection.",
        "comments": f"Test Data: source_agent=agent_cx_smoke_01; target_agent=agent_cx_smoke_02; target_queue=Escalations; call_start={fmt_dt(CALL_START + timedelta(minutes=10))}.",
    },
    {
        "scenario": "NICE Smoke - Keypad Availability and Function",
        "description": "Verify the keypad is available during a call and can send DTMF input correctly when needed.",
        "preconditions": "Active call connected; call flow or test IVR available to accept keypad input.",
        "steps": "1. Connect a call that supports keypad input\n2. Open the keypad during the call\n3. Press test digits\n4. Observe whether the system accepts the DTMF input correctly",
        "expected": "The keypad opens during the call and the entered digits are sent successfully to the connected system or IVR.",
        "comments": f"Test Data: ivr_test_number=+16145550170; dtmf_sequence=1234#; call_start={fmt_dt(CALL_START + timedelta(minutes=20))}.",
    },
    {
        "scenario": "NICE Smoke - End Call Function",
        "description": "Verify the agent can end an active call from NICE CXone and the call closes cleanly.",
        "preconditions": "Agent is connected to an active NICE CXone call.",
        "steps": "1. Accept or place a test call\n2. Select the end-call option\n3. Observe call status after ending\n4. Confirm the interaction closes correctly",
        "expected": "The call ends successfully, the status changes to completed or closed, and the agent returns to the expected post-call or available state.",
        "comments": f"Test Data: active_call_id=SMOKE-CALL-06; call_start={fmt_dt(CALL_START + timedelta(minutes=30))}; expected_end_time={fmt_dt(CALL_END + timedelta(minutes=30))}.",
    },
    {
        "scenario": "NICE Smoke - Directory Access During Call",
        "description": "Verify the agent can open and use the directory while on an active call in NICE CXone.",
        "preconditions": "Agent is on an active call and directory access is enabled.",
        "steps": "1. Accept or place a test call\n2. Open the directory during the call\n3. Search for an agent, queue, or contact\n4. Confirm the directory remains usable while the call stays active",
        "expected": "The directory opens and can be used during the call without interrupting the active interaction.",
        "comments": f"Test Data: directory_search_value=agent_cx_smoke_02; call_start={fmt_dt(CALL_START + timedelta(minutes=40))}; expected_call_state=Connected.",
    },
    {
        "scenario": "NICE Smoke - Help Center Access During Call",
        "description": "Verify the help center can be opened while the agent is on a call and does not break the active call session.",
        "preconditions": "Agent is on an active call; help center link or icon is available in NICE CXone.",
        "steps": "1. Accept or place a test call\n2. Open the help center during the call\n3. Confirm help content loads\n4. Verify the active call remains connected and usable",
        "expected": "The help center opens successfully during the call, and the active call remains stable and connected.",
        "comments": f"Test Data: help_center_entry=Voice Controls Help; call_start={fmt_dt(CALL_START + timedelta(minutes=50))}; expected_call_state=Connected.",
    },
    {
        "scenario": "NICE Smoke - Agent Availability State Changes",
        "description": "Verify the agent can change status between Available, Not Available, and Away in NICE CXone.",
        "preconditions": "Agent is logged in successfully to NICE CXone and no active call is blocking manual state change.",
        "steps": "1. Change agent state to Available\n2. Change agent state to Not Available\n3. Change agent state to Away\n4. Confirm each state change is reflected in the NICE CXone UI",
        "expected": "The agent can change between Available, Not Available, and Away, and each selected state is reflected correctly.",
        "comments": f"Test Data: agent_id=agent_cx_smoke_01; states=Available,Not Available,Away; validation_time={fmt_dt(CALL_START + timedelta(hours=1))}.",
    },
]


def apply_row_style(worksheet, target_row: int, style_row: int) -> None:
    for column in range(1, worksheet.max_column + 1):
        source_cell = worksheet.cell(style_row, column)
        target_cell = worksheet.cell(target_row, column)
        if source_cell.has_style:
            target_cell._style = copy(source_cell._style)
        if source_cell.font:
            target_cell.font = copy(source_cell.font)
        if source_cell.fill:
            target_cell.fill = copy(source_cell.fill)
        if source_cell.border:
            target_cell.border = copy(source_cell.border)
        if source_cell.alignment:
            target_cell.alignment = copy(source_cell.alignment)
        else:
            target_cell.alignment = Alignment(wrap_text=True, vertical="top")
        if source_cell.number_format:
            target_cell.number_format = source_cell.number_format
        if source_cell.protection:
            target_cell.protection = copy(source_cell.protection)


def main() -> None:
    workbook = load_workbook(TEMPLATE_PATH)
    worksheet = workbook.active

    if worksheet.max_row > 1:
        worksheet.delete_rows(2, worksheet.max_row - 1)

    style_row = 2
    worksheet.insert_rows(2, amount=len(TEST_CASES))

    for index, test_case in enumerate(TEST_CASES, start=2):
        apply_row_style(worksheet, index, style_row + len(TEST_CASES))
        worksheet.cell(index, 1).value = index - 1
        worksheet.cell(index, 2).value = test_case["scenario"]
        worksheet.cell(index, 3).value = test_case["description"]
        worksheet.cell(index, 4).value = test_case["preconditions"]
        worksheet.cell(index, 5).value = test_case["steps"]
        worksheet.cell(index, 6).value = test_case["expected"]
        worksheet.cell(index, 7).value = None
        worksheet.cell(index, 8).value = None
        worksheet.cell(index, 9).value = test_case["comments"]
        for column in range(1, 10):
            worksheet.cell(index, column).alignment = Alignment(
                wrap_text=True,
                vertical="top",
            )

    workbook.save(OUTPUT_PATH)
    print(OUTPUT_PATH)


if __name__ == "__main__":
    main()